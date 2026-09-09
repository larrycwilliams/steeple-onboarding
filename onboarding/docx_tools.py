"""Run-aware text replacement inside .docx files.

Word splits a single sentence across many <w:r> runs (spell-check state,
formatting, rsid churn). A naive paragraph.text = ... rewrite destroys the
bold/colour runs that make these documents look designed. Everything here
edits runs in place so the original formatting survives.
"""
from __future__ import annotations

import re
from docx.document import Document as _Doc
from docx.table import Table
from docx.text.paragraph import Paragraph


def iter_block_paragraphs(container):
    """Yield every paragraph in a document part, including inside tables."""
    if isinstance(container, _Doc):
        body = container.element.body
        parent = container
    else:
        body = container._element
        parent = container

    for para in parent.paragraphs:
        yield para
    for table in parent.tables:
        yield from _iter_table_paragraphs(table)


def _iter_table_paragraphs(table: Table):
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                yield para
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def all_paragraphs(doc: _Doc):
    """Every paragraph in the document body, headers and footers."""
    yield from iter_block_paragraphs(doc)
    for section in doc.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        ):
            if part is None:
                continue
            try:
                yield from iter_block_paragraphs(part)
            except Exception:
                continue


def replace_once(para: Paragraph, old: str, new: str) -> bool:
    """Replace the first occurrence of `old` in `para`, preserving run formatting.

    The replacement text lands entirely in the first run that the match
    touches, so it inherits that run's formatting. Returns False if not found.
    """
    runs = para.runs
    if not runs:
        return False

    spans = []
    pos = 0
    for run in runs:
        length = len(run.text)
        spans.append((pos, pos + length))
        pos += length

    full = "".join(run.text for run in runs)
    idx = full.find(old)
    if idx < 0:
        return False
    end = idx + len(old)

    written = False
    for run, (start, stop) in zip(runs, spans):
        if stop <= idx or start >= end:
            continue
        local_start = max(idx, start) - start
        local_end = min(end, stop) - start
        text = run.text
        if not written:
            run.text = text[:local_start] + new + text[local_end:]
            written = True
        else:
            run.text = text[:local_start] + text[local_end:]
    return written


def replace_all(para: Paragraph, old: str, new: str, limit: int = 40) -> int:
    count = 0
    while count < limit and replace_once(para, old, new):
        count += 1
    return count


def apply_map(doc: _Doc, mapping: dict[str, str]) -> dict[str, int]:
    """Apply a literal -> replacement map across the whole document.

    Two passes via sentinels so a replacement containing its own search text
    can never loop, and so longer keys always win over shorter ones.
    """
    keys = sorted(mapping, key=len, reverse=True)
    sentinels = {key: f"@@SS{i:03d}@@" for i, key in enumerate(keys)}
    hits = {key: 0 for key in keys}

    paragraphs = list(all_paragraphs(doc))

    for key in keys:
        token = sentinels[key]
        for para in paragraphs:
            hits[key] += replace_all(para, key, token)

    for key in keys:
        token = sentinels[key]
        value = mapping[key]
        for para in paragraphs:
            replace_all(para, token, value)

    return hits


def document_text(doc: _Doc) -> str:
    return "\n".join(p.text for p in all_paragraphs(doc))


def find_unreplaced(doc: _Doc, needles: list[str]) -> list[str]:
    """Report any source-specific strings that survived templatizing."""
    text = document_text(doc)
    return [n for n in needles if n in text]


BLANK_RE = re.compile(r"_{6,}")


def blank_fields(doc: _Doc) -> list[str]:
    """Paragraphs that still contain fill-in-the-blank underscores."""
    out = []
    for para in all_paragraphs(doc):
        if BLANK_RE.search(para.text):
            out.append(para.text.strip())
    return out
